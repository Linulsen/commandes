#!/usr/bin/env python3
"""Importe le classeur « Aide a la commande » dans Supabase.

Le classeur porte deux familles d'onglets qui ne se lisent pas de la même façon :

* les six chambres livrées par GLD, où la colonne E porte le multiplicateur réel
  (unités livrées par colis) et où chaque semaine occupe quatre colonnes
  Stock / Conso / Cde / Total ;
* les fournisseurs à liste propre (France Boissons, Cledor), où ce multiplicateur
  s'appelle « Fact » en colonne C, et où le premier bloc peut n'avoir que trois
  colonnes.

Le script ne se fie donc pas à des décalages fixes : il relit les en-têtes sous
chaque date pour savoir quelle colonne porte quoi.

Usage :  SUPABASE_URL=… SUPABASE_SERVICE_ROLE_KEY=… python3 scripts/import_excel.py [chemin.xlsx]
"""
import datetime
import os
import re
import sys
import unicodedata

import openpyxl
import requests

XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "..",
    "Aide a la commande - 22.06.2026.xlsx")
URL = os.environ["SUPABASE_URL"].rstrip("/")
KEY = os.environ["SUPABASE_SERVICE_ROLE_KEY"]

S = requests.Session()
S.headers.update({
    "apikey": KEY,
    "Authorization": f"Bearer {KEY}",
    "Content-Type": "application/json",
})


def rest(method, table, payload=None, params=None, prefer=None, on_conflict=None):
    headers = {"Prefer": prefer} if prefer else {}
    if on_conflict:
        params = dict(params or {}, on_conflict=on_conflict)
    r = S.request(method, f"{URL}/rest/v1/{table}", json=payload,
                  params=params, headers=headers, timeout=120)
    if not r.ok:
        raise SystemExit(f"{method} {table} -> {r.status_code} {r.text[:600]}")
    return r.json() if r.text.strip() else []


def slug(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"-+", "-", re.sub(r"[^a-z0-9]+", "-", s.lower())).strip("-")


def num(v):
    if v is None or isinstance(v, str) and not v.strip():
        return None
    try:
        return round(float(v), 4)
    except (TypeError, ValueError):
        return None


def as_date(v):
    """Les dates du classeur sont parfois saisies à la main (« 01/072026 »)."""
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{1,2})\s*/\s*(\d{1,2})\s*/?\s*(\d{4})\s*$", v)
        if m:
            d, mo, y = (int(x) for x in m.groups())
            try:
                return datetime.date(y, mo, d)
            except ValueError:
                return None
    return None


# --- Découpage du classeur -------------------------------------------------
# nom d'onglet -> (fournisseur, libellé de la zone, secteur, fact en colonne E ?)
ONGLETS = [
    ("Cave",                   "GLD", "Cave",                     None,                     True),
    ("Chambre froide boisson", "GLD", "Chambre froide boisson",   None,                     True),
    ("Reserve seche",          "GLD", "Réserve sèche",            "GLD - Ambiant",          True),
    ("Economat",               "GLD", "Economat",                 "GLD - Economat/Épice",   True),
    ("Chambre froide",         "GLD", "Chambre froide",           "GLD - Frais",            True),
    ("Chambre negative",       "GLD", "Chambre négative",         "GLD - Surgelés",         True),
    ("France Boissons",        "France Boissons", "Boissons",     None,                     False),
    ("Cledor",                 "Cledor", "Fruits et légumes",     None,                     False),
]
FOURNISSEURS = [
    ("GLD",             "hebdo",    "Commande hebdomadaire, une passe par chambre", 1),
    ("France Boissons", "hebdo",    None,                                           2),
    ("Cledor",          "bi-hebdo", "Dimanche pour mardi, mercredi pour vendredi",  3),
]

ROLES = {"stock": "stock", "cde": "colis", "conso": "conso", "total": "total"}


def role(header):
    h = slug(header or "")
    if h.startswith("stock"):
        return "stock"
    if h.startswith("conso"):
        return "conso"
    if h.startswith("cde"):
        return "cde"
    if h.startswith("total"):
        return "total"
    return None


def lire_onglet(ws):
    """Retourne (produits, blocs) où un bloc = (date, libellé, {rôle: colonne}).

    La largeur balayée suit max_column : Cledor, relevé deux fois par semaine,
    pousse ses blocs jusqu'à la colonne 108 là où une chambre s'arrête à 62.
    """
    large = min(max(ws.max_column, 8) + 8, 220)  # max_column part en vrille
    # sur les onglets au formatage étalé (16 000 colonnes vides).
    hrow = next(r for r in range(1, 9)
                if slug(ws.cell(r, 1).value or "") == "produit")
    entetes = {c: ws.cell(hrow, c).value for c in range(1, large)}

    # La ligne des dates est la dernière, au-dessus des en-têtes, qui en porte.
    drow = max((r for r in range(1, hrow)
                if any(as_date(ws.cell(r, c).value) for c in range(1, large))),
               default=None)
    if drow is None:
        return [], []
    # Le libellé de rythme (Cledor) vit sur une ligne au-dessus des dates.
    lrow = next((r for r in range(1, drow)
                 if any(isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip()
                        for c in range(4, 80))), None)

    # Un bloc commence à sa colonne « Stock » : c'est le seul repère présent
    # partout. La date, elle, manque parfois (Cledor, semaine du 19/06) et les
    # onglets traînent des blocs de gabarit vides après le dernier relevé.
    departs = [c for c in range(1, large) if role(entetes.get(c)) == "stock"]
    produits_rows = [r for r in range(hrow + 1, ws.max_row + 1)
                     if ws.cell(r, 1).value and str(ws.cell(r, 1).value).strip()]

    bruts = []
    for i, c in enumerate(departs):
        fin = departs[i + 1] if i + 1 < len(departs) else large
        cols = {}
        for cc in range(c, min(fin, c + 6)):
            r_ = role(entetes.get(cc))
            if r_ and r_ not in cols:
                cols[r_] = cc
        # Un bloc du classeur vaut toujours Stock / Conso / Cde / Total, mais
        # l'en-tête « Cde » manque parfois (France Boissons, semaine du 21/06)
        # alors que la colonne, elle, porte bien des quantités. La colonne juste
        # avant le total est la commande : c'est ce que dit la formule du total.
        if "cde" not in cols and "total" in cols and cols["total"] - 1 > c:
            candidate = cols["total"] - 1
            if candidate not in cols.values():
                cols["cde"] = candidate
        d = as_date(ws.cell(drow, c).value)
        v = ws.cell(lrow, c).value if lrow else None
        # « dimanche pour mardi », « Dimanche pour Mardi »... : un seul libellé.
        lib = " ".join(str(v).split()).capitalize() if isinstance(v, str) else ""
        # Un relevé commence toujours par un comptage. Une colonne de gabarit
        # peut traîner une quantité commandée saisie d'avance : sans stock
        # relevé, ce n'est pas une campagne, et en faire une session créerait
        # un trou dans la chaîne des consommations.
        releve = any(num(ws.cell(r, cols["stock"]).value) is not None
                     for r in produits_rows)
        if not releve:
            continue
        bruts.append([d, lib, cols])

    # Une date absente sur un bloc qui porte des relevés est reconstituée au
    # milieu de ses voisins : la date exacte importe moins que l'ordre, dont
    # dépend le calcul de la consommation d'un relevé au suivant.
    for i, b in enumerate(bruts):
        if b[0] is not None:
            continue
        avant = next((bruts[j][0] for j in range(i - 1, -1, -1) if bruts[j][0]), None)
        apres = next((bruts[j][0] for j in range(i + 1, len(bruts)) if bruts[j][0]), None)
        if avant and apres:
            b[0] = avant + datetime.timedelta(days=(apres - avant).days // 2)
        elif avant:
            b[0] = avant + datetime.timedelta(days=3 if len(bruts) > 20 else 7)
        elif apres:
            b[0] = apres - datetime.timedelta(days=3 if len(bruts) > 20 else 7)
    blocs = [(b[0], b[1], b[2]) for b in bruts if b[0]]

    produits = [(r, str(ws.cell(r, 1).value).strip()) for r in produits_rows]
    return produits, blocs


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    if "--reset" in sys.argv:
        # Relire le classeur doit pouvoir repartir d'une base propre : les
        # dates reconstituées changent si le classeur évolue, et une session
        # orpheline fausserait la chaîne des consommations.
        for t in ("cmd_lignes", "cmd_sessions"):
            rest("DELETE", t, params={"id": "gt.0"})
        print("relevés précédents effacés")

    rest("POST", "cmd_fournisseurs",
         [{"nom": n, "slug": slug(n), "frequence": f, "rythme": ry, "ordre": o}
          for n, f, ry, o in FOURNISSEURS],
         prefer="resolution=merge-duplicates", on_conflict="slug")
    fourn = {f["nom"]: f["id"] for f in rest("GET", "cmd_fournisseurs",
                                            params={"select": "id,nom"})}

    zones_payload = [
        {"fournisseur_id": fourn[f], "nom": nom, "slug": slug(f"{f}-{nom}"),
         "secteur": sect, "ordre": i}
        for i, (onglet, f, nom, sect, _) in enumerate(ONGLETS)
    ]
    rest("POST", "cmd_zones", zones_payload,
         prefer="resolution=merge-duplicates", on_conflict="slug")
    zones = {z["slug"]: z["id"] for z in rest("GET", "cmd_zones",
                                             params={"select": "id,slug"})}

    # --- Produits ---------------------------------------------------------
    produits_payload, lecture = [], {}
    for onglet, f, nom, sect, fact_en_e in ONGLETS:
        ws = wb[onglet]
        produits, blocs = lire_onglet(ws)
        lecture[onglet] = (ws, produits, blocs, f, fact_en_e)
        zid = zones[slug(f"{f}-{nom}")]
        for i, (r, label) in enumerate(produits):
            cond = ws.cell(r, 2).value
            upc = num(ws.cell(r, 3).value)
            unite = ws.cell(r, 4).value if fact_en_e else None
            fact = num(ws.cell(r, 5).value) if fact_en_e else num(ws.cell(r, 3).value)
            produits_payload.append({
                "zone_id": zid, "nom": label,
                "conditionnement": str(cond).strip() if cond else None,
                "unites_par_conditionnement": upc,
                "unite": str(unite).strip() if unite else None,
                "fact": fact or upc or 1,
                "ordre": i,
            })
    rest("POST", "cmd_produits", produits_payload,
         prefer="resolution=merge-duplicates", on_conflict="zone_id,nom")
    pid = {(p["zone_id"], p["nom"]): p["id"]
           for p in rest("GET", "cmd_produits",
                         params={"select": "id,zone_id,nom", "limit": "5000"})}
    print(f"{len(produits_payload)} produits")

    # --- Sessions ---------------------------------------------------------
    # Les six chambres GLD partagent un calendrier, à quelques saisies près
    # (un 21/06 pour un 22/06). On aligne sur le calendrier de la Réserve sèche
    # pour qu'une semaine reste une seule commande.
    canon = sorted({d for d, _, _ in lecture["Reserve seche"][2]})

    def aligner(d, fournisseur):
        if fournisseur != "GLD":
            return d
        proche = min(canon, key=lambda c: abs((c - d).days))
        return proche if abs((proche - d).days) <= 3 else d

    sessions_payload = {}
    for onglet, (ws, produits, blocs, f, _) in lecture.items():
        for d, lib, _cols in blocs:
            key = (fourn[f], aligner(d, f).isoformat(), lib)
            sessions_payload[key] = {
                "fournisseur_id": key[0], "date_commande": key[1],
                "libelle": key[2], "statut": "validee",
                "note": "Repris du classeur Excel",
            }
    rest("POST", "cmd_sessions", list(sessions_payload.values()),
         prefer="resolution=merge-duplicates",
         on_conflict="fournisseur_id,date_commande,libelle")
    sid = {(s["fournisseur_id"], s["date_commande"], s["libelle"]): s["id"]
           for s in rest("GET", "cmd_sessions",
                         params={"select": "id,fournisseur_id,date_commande,libelle",
                                 "limit": "2000"})}
    print(f"{len(sessions_payload)} sessions")

    # --- Lignes -----------------------------------------------------------
    lignes = []
    for onglet, f, nom, sect, _ in ONGLETS:
        ws, produits, blocs, _f, _e = lecture[onglet]
        zid = zones[slug(f"{f}-{nom}")]
        for i, (d, lib, cols) in enumerate(blocs):
            s = sid[(fourn[f], aligner(d, f).isoformat(), lib)]
            for r, label in produits:
                stock = num(ws.cell(r, cols["stock"]).value)
                colis = num(ws.cell(r, cols["cde"]).value) if "cde" in cols else None
                # La consommation se déduit du relevé précédent ; seul le tout
                # premier bloc doit la porter en dur, sinon elle serait comptée
                # deux fois.
                conso_m = (num(ws.cell(r, cols["conso"]).value)
                           if i == 0 and "conso" in cols else None)
                if stock is None and colis is None and conso_m is None:
                    continue
                lignes.append({"session_id": s, "produit_id": pid[(zid, label)],
                               "stock": stock, "colis": colis,
                               "conso_manuelle": conso_m})

    for i in range(0, len(lignes), 500):
        rest("POST", "cmd_lignes", lignes[i:i + 500],
             prefer="resolution=merge-duplicates",
             on_conflict="session_id,produit_id")
    print(f"{len(lignes)} lignes de relevé")


if __name__ == "__main__":
    main()
